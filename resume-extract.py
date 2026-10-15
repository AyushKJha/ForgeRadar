"""Local-only, bounded résumé text extraction. Prints suggestions, never the full file."""
import json
import re
import sys
import zipfile
from pathlib import Path


def clean(value):
    return re.sub(r"\s+", " ", value).strip()


def extract_pdf(filename):
    from pypdf import PdfReader

    reader = PdfReader(filename, strict=True)
    if reader.is_encrypted or len(reader.pages) > 20:
        raise ValueError("Encrypted or overlong PDF documents are unsupported")
    root = reader.trailer.get("/Root", {})
    if any(key in root for key in ("/OpenAction", "/AA", "/Names")):
        raise ValueError("PDF active content is unsupported")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_docx(filename):
    from docx import Document

    with zipfile.ZipFile(filename) as archive:
        infos = archive.infolist()
        if len(infos) > 1000 or sum(item.file_size for item in infos) > 12_000_000:
            raise ValueError("DOCX archive is too large")
        for item in infos:
            parts = item.filename.replace("\\", "/").split("/")
            if ".." in parts or item.filename.startswith("/") or "vbaproject" in item.filename.lower():
                raise ValueError("DOCX contains unsupported paths or active content")
    document = Document(filename)
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def suggestions(text):
    patterns = {
        "languages": {
            "Python": r"\bpython\b", "JavaScript": r"\bjavascript\b", "TypeScript": r"\btypescript\b",
            "C++": r"\bc\+\+\b", "Java": r"\bjava\b", "Rust": r"\brust\b",
        },
        "domains": {
            "cybersecurity": r"\b(cybersecurity|cyber security|penetration testing|threat detection)\b",
            "embedded_systems": r"\b(embedded systems?|microcontrollers?|firmware)\b",
            "signal_processing": r"\b(signal processing|digital signal processing|dsp)\b",
            "ai_ml": r"\b(machine learning|artificial intelligence|\bAI\b)\b",
            "cloud": r"\b(cloud computing|aws|azure|gcp)\b",
            "robotics": r"\brobotics?\b",
        },
        "frameworks": {
            "React": r"\breact(?:\.js)?\b", "PyTorch": r"\bpytorch\b", "TensorFlow": r"\btensorflow\b",
            "FastAPI": r"\bfastapi\b", "Django": r"\bdjango\b", "Node.js": r"\bnode(?:\.js)?\b",
        },
    }
    lines = [clean(line) for line in text.splitlines() if clean(line)]
    found = {"languages": [], "domains": [], "frameworks": [], "education": [], "projects": [], "experience": [], "researchInterests": [], "goals": []}
    for category, terms in patterns.items():
        for label, pattern in terms.items():
            for line in lines:
                if re.search(pattern, line, re.I):
                    found[category].append({"value": label, "evidence": line[:180], "verified": False})
                    break
    line_patterns = {
        "education": r"\b(bachelor|master|b\.?tech|m\.?tech|ph\.?d|university|college|degree)\b",
        "projects": r"\b(project|built|implemented|developed|prototype)\b",
        "experience": r"\b(intern|research assistant|engineer|employment|work experience)\b",
        "researchInterests": r"\b(research interests?|research areas?|interested in)\b",
        "goals": r"\b(career goal|objective|aspir(?:e|ing)|looking for)\b",
    }
    for category, pattern in line_patterns.items():
        for line in lines:
            if re.search(pattern, line, re.I):
                found[category].append({"value": line[:180], "evidence": line[:180], "verified": False})
                if len(found[category]) >= 5:
                    break
    return found


def main():
    filename = Path(sys.argv[1])
    if filename.stat().st_size > 6_000_000:
        raise ValueError("Résumé exceeds the 6 MB limit")
    text = extract_pdf(filename) if filename.suffix.lower() == ".pdf" else extract_docx(filename)
    text = text[:20000]
    if not text.strip():
        raise ValueError("No readable text found; scanned images require manual entry")
    print(json.dumps({"preview": clean(text[:2400]), "suggestions": suggestions(text), "note": "Suggestions are claims from the document, not verified skills."}))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error)[:250]}))
        sys.exit(1)
